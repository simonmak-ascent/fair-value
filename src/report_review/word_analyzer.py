"""
Word Document Analyzer Module
"""

import re
from typing import Dict, List
import logging

logger = logging.getLogger(__name__)

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def extract_text(file_path: str) -> str:
    """
    Extract text from Word document
    
    Args:
        file_path: Path to Word file
    
    Returns:
        Text content
    """
    if not DOCX_AVAILABLE:
        return ""
    
    try:
        doc = docx.Document(file_path)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        return text
    except Exception as e:
        logger.error(f"Error extracting Word text: {e}")
        return ""


def extract_tables(file_path: str) -> List[List[List[str]]]:
    """
    Extract tables from Word document
    
    Args:
        file_path: Path to Word file
    
    Returns:
        List of tables
    """
    if not DOCX_AVAILABLE:
        return []
    
    try:
        doc = docx.Document(file_path)
        tables = []
        
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                table_data.append(row_data)
            tables.append(table_data)
        
        return tables
    except Exception as e:
        logger.error(f"Error extracting tables: {e}")
        return []


def extract_valuation_conclusion(text: str) -> Dict:
    """
    Extract valuation conclusion
    
    Args:
        text: Document text
    
    Returns:
        Valuation conclusion
    """
    fair_value = re.search(
        r'fair\s+value[:\s]+[\$€£¥]?\s*([\d,]+(?:\.\d+)?)',
        text, re.IGNORECASE
    )
    
    methodology = re.search(
        r'(DCF|NAV|CCA|multiples|dcf|nav)',
        text, re.IGNORECASE
    )
    
    return {
        'fair_value': fair_value.group(1) if fair_value else None,
        'methodology': methodology.group(1) if methodology else None
    }


def check_ivs_compliance(text: str) -> Dict:
    """
    Check IVS/IFRS compliance
    
    Args:
        text: Document text
    
    Returns:
        Compliance information
    """
    standards = []
    
    if re.search(r'IVS|valuation\s+standard', text, re.IGNORECASE):
        standards.append('IVS')
    
    if re.search(r'IFRS|IAS', text, re.IGNORECASE):
        standards.append('IFRS/IAS')
    
    return {
        'has_standards': len(standards) > 0,
        'standards': standards
    }


def analyze_word_report(file_path: str) -> Dict:
    """
    Main entry point for Word document analysis
    
    Args:
        file_path: Path to Word file
    
    Returns:
        Analysis results
    """
    try:
        text = extract_text(file_path)
        tables = extract_tables(file_path)
        
        if not text:
            return {
                'file': file_path,
                'status': 'error',
                'message': 'Could not extract text from document'
            }
        
        return {
            'file': file_path,
            'status': 'success',
            'text_length': len(text),
            'n_tables': len(tables),
            'valuation': extract_valuation_conclusion(text),
            'compliance': check_ivs_compliance(text)
        }
    except Exception as e:
        logger.error(f"Error analyzing Word document: {e}")
        return {
            'file': file_path,
            'status': 'error',
            'message': str(e)
        }
